"""Load and normalize the invoice Excel, the customer master and the organization profile."""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime

from openpyxl import load_workbook

FX_TO_SAR = {"SAR": 1.0, "USD": 3.75}  # USD uses the fixed SAMA peg; other currencies are rejected

# Column synonyms so differently-labelled exports still map to the same fields.
COLUMNS = {
    "number": ["invoice no", "invoice number", "inv no", "رقم الفاتورة"],
    "issue_date": ["invoice date", "date", "تاريخ الفاتورة"],
    "due_date": ["due date", "تاريخ الاستحقاق"],
    "customer": ["customer name", "customer", "اسم العميل"],
    "customer_vat": ["customer vat no", "customer vat", "الرقم الضريبي للعميل"],
    "seller_vat": ["seller vat no", "seller vat"],
    "kind": ["document type", "type"],
    "ref": ["reference invoice", "reference"],
    "currency": ["currency", "العملة"],
    "net": ["amount (excl. vat)", "amount excl vat", "net amount", "المبلغ قبل الضريبة"],
    "vat": ["vat 15%", "vat", "الضريبة"],
    "total": ["total", "total incl vat", "الإجمالي"],
}


def _parse_date(v) -> date | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _f(v) -> float:
    if v is None or v == "":
        return 0.0
    return float(str(v).replace(",", ""))


def load_invoices(source, org_vat: str | None) -> dict:
    """Returns {"invoices": [...], "issues": [...], "stats": {...}}. Amounts are in SAR."""
    data = source if isinstance(source, (bytes, bytearray)) else open(source, "rb").read()
    ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))

    # Find the header row instead of assuming row 1 (exports often start with a title block).
    header_idx, colmap = None, {}
    for i, r in enumerate(rows[:20]):
        labels = [str(c).strip().lower() if c is not None else "" for c in r]
        found = {k: labels.index(s) for k, syns in COLUMNS.items() for s in syns if s in labels}
        if "number" in found and "net" in found:
            header_idx, colmap = i, found
            break
    if header_idx is None:
        raise ValueError("Could not find the header row. Expected columns like 'Invoice No' and 'Amount (excl. VAT)'.")

    invoices, issues, seen = [], [], {}
    stats = {"rows_read": 0, "duplicates_removed": 0, "credit_notes": 0, "fx_converted": 0,
             "entity_mismatch": 0, "rejected": 0}
    for excel_row, r in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        get = lambda k: r[colmap[k]] if k in colmap and colmap[k] < len(r) else None
        if not get("number"):
            continue
        stats["rows_read"] += 1
        number = str(get("number")).strip()
        issue = _parse_date(get("issue_date"))
        currency = (str(get("currency") or "SAR")).strip().upper()
        source_ref = f"{ws.title}!row {excel_row}"
        if issue is None:
            issues.append({"type": "bad_date", "invoice": number, "source": source_ref,
                           "message": f"Invoice {number}: date not recognised. Use DD/MM/YYYY."})
            stats["rejected"] += 1
            continue
        if currency not in FX_TO_SAR:
            issues.append({"type": "unsupported_currency", "invoice": number, "source": source_ref,
                           "message": f"Invoice {number} is in {currency}. Only SAR and USD are supported."})
            stats["rejected"] += 1
            continue
        fx = FX_TO_SAR[currency]
        net, vat = _f(get("net")), _f(get("vat"))
        total = _f(get("total")) or net + vat
        kind = "credit_note" if "credit" in str(get("kind") or "").lower() or net < 0 else "invoice"
        customer = str(get("customer") or "").strip()
        inv = {
            "number": number,
            "kind": kind,
            "issue_date": issue.isoformat(),
            "due_date": (_parse_date(get("due_date")) or issue).isoformat(),
            "customer_name": customer,
            "customer_vat": str(get("customer_vat") or "").strip(),
            "seller_vat": str(get("seller_vat") or "").strip(),
            "ref": str(get("ref") or "").strip(),
            "original_currency": currency,
            "original_total": round(total, 2),
            "fx_rate": fx,
            "net_sar": round(net * fx, 2),
            "vat_sar": round(vat * fx, 2),
            "gross_sar": round(total * fx, 2),
            "source_ref": source_ref,
            "flags": [],
        }
        key = (number, round(total, 2), customer)
        if key in seen:
            stats["duplicates_removed"] += 1
            issues.append({"type": "duplicate", "invoice": number, "source": source_ref,
                           "message": f"Invoice {number} appears twice (also {seen[key]}). The duplicate was ignored."})
            continue
        if number in {i["number"] for i in invoices[-500:]}:
            inv["flags"].append("number_reused")
        seen[key] = source_ref
        if currency != "SAR":
            stats["fx_converted"] += 1
            inv["flags"].append("fx_converted")
        if org_vat and inv["seller_vat"] and inv["seller_vat"] != org_vat:
            stats["entity_mismatch"] += 1
            inv["flags"].append("entity_mismatch")
        if kind == "credit_note":
            stats["credit_notes"] += 1
        invoices.append(inv)

    # Net credit notes against the invoice they reference.
    by_no = {i["number"]: i for i in invoices if i["kind"] == "invoice"}
    for cn in (i for i in invoices if i["kind"] == "credit_note"):
        target = by_no.get(cn["ref"])
        if target:
            target["credit_note"] = cn["number"]
            target["credit_note_gross_sar"] = cn["gross_sar"]
        else:
            issues.append({"type": "orphan_credit_note", "invoice": cn["number"], "source": cn["source_ref"],
                           "message": f"Credit note {cn['number']} references {cn['ref'] or 'no invoice'}, which was not found."})
    return {"invoices": invoices, "issues": issues, "stats": stats}


def load_customers(source) -> list[dict]:
    if source is None:
        return []
    text = source.decode("utf-8-sig") if isinstance(source, (bytes, bytearray)) else open(source, encoding="utf-8-sig").read()
    return list(csv.DictReader(io.StringIO(text)))


def load_org(source) -> dict:
    if source is None:
        return {}
    text = source.decode("utf-8") if isinstance(source, (bytes, bytearray)) else open(source, encoding="utf-8").read()
    return json.loads(text)


def norm_ar(s: str) -> str:
    s = re.sub("[إأآا]", "ا", s or "")
    s = s.replace("ة", "ه").replace("ى", "ي")
    return " ".join(s.split())
